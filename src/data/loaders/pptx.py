from .base import BaseLoader
from pptx import Presentation
from src.core.exceptions import FileReadError, EmptyDocumentError
from src.core.logger import get_logger

logger = get_logger("data.loaders.pptx")


class PPTXLoader(BaseLoader):

    def load(self, file_path: str) -> str:
        logger.info(f"Loading PPTX file: {file_path}")
        try:
            prs = Presentation(file_path)
            text_parts = []

            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)

            result = "\n".join(text_parts)
            logger.debug("Extracted text from PPTX slides")

        except Exception as e:
            logger.exception(f"Error reading PPTX file: {file_path}")
            raise FileReadError(str(file_path), str(e))

        if not result.strip():
            logger.warning(f"Empty PPTX document: {file_path}")
            raise EmptyDocumentError(file_path)

        logger.info(f"Successfully loaded PPTX file: {file_path} ({len(result)} chars)")

        return result
