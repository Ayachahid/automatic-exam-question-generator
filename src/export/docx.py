from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .base import BaseExporter
from src.core.logger import get_logger

logger = get_logger("export.docx")

# ── colour palette (matches project branding) ─────────────────────────────────
COLOR_TITLE = RGBColor(0x1F, 0x49, 0x7D)  # dark blue  – document title
COLOR_Q_NUM = RGBColor(0x2E, 0x75, 0xB6)  # mid  blue  – "Question N"
COLOR_LABEL = RGBColor(0x70, 0x70, 0x70)  # grey       – "Answer:", "Options:"


def _add_colored_run(paragraph, text: str, bold: bool, color: RGBColor, size_pt: int):
    """Helper: add a styled run to an existing paragraph."""
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    return run


class DOCXExporter(BaseExporter):
    """
    Exports questions to a Microsoft Word (.docx) document.

    Handles all question types:
        - Multiple Choice  (options list + letter answer)
        - True / False
        - Short Answer
        - Essay
        - Scenario-Based   (scenario block + concepts)
    """

    def export(self, questions: list[dict], output_path: str) -> str:
        """
        Build and save a .docx file from a list of question dicts.

        Args:
            questions:    List of question dicts (pipeline output).
            output_path:  Destination .docx file path.

        Returns:
            Absolute path to the exported file.
        """
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        self._set_document_style(doc)
        self._add_title(doc, total=len(questions))

        for idx, q in enumerate(questions, start=1):
            self._add_question_block(doc, idx, q)
            doc.add_paragraph()  # spacer between questions

        doc.save(str(path))
        logger.info(f"[DOCXExporter] Exported {len(questions)} questions → {path}")
        return str(path.absolute())

    # ── private helpers ────────────────────────────────────────────────────────

    def _set_document_style(self, doc: Document):
        """Apply base font settings to the Normal style."""
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

    def _add_title(self, doc: Document, total: int):
        """Insert the document header."""
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_colored_run(
            title_para,
            "📝  Exam Question Bank",
            bold=True,
            color=COLOR_TITLE,
            size_pt=18,
        )
        doc.add_paragraph(f"Total questions: {total}").runs[
            0
        ].font.color.rgb = COLOR_LABEL
        doc.add_paragraph()  # spacer

    def _add_question_block(self, doc: Document, idx: int, q: dict):
        """Render a single question with all its fields."""
        # ── Question number + text ─────────────────────────────────────────
        q_para = doc.add_paragraph()
        _add_colored_run(
            q_para, f"Question {idx}:  ", bold=True, color=COLOR_Q_NUM, size_pt=12
        )
        q_para.add_run(q.get("question", "")).font.size = Pt(12)

        # ── Scenario block (scenario-based questions) ──────────────────────
        if q.get("scenario"):
            sc_para = doc.add_paragraph()
            _add_colored_run(
                sc_para, "Scenario:  ", bold=True, color=COLOR_LABEL, size_pt=11
            )
            sc_para.add_run(q["scenario"])

        # ── Options (MCQ) ──────────────────────────────────────────────────
        if q.get("options"):
            opt_para = doc.add_paragraph()
            _add_colored_run(
                opt_para, "Options:", bold=True, color=COLOR_LABEL, size_pt=11
            )
            letters = ["A", "B", "C", "D", "E"]
            for i, option in enumerate(q["options"]):
                letter = letters[i] if i < len(letters) else str(i + 1)
                doc.add_paragraph(f"    {letter}.  {option}", style="List Bullet")

        # ── Answer ─────────────────────────────────────────────────────────
        ans_para = doc.add_paragraph()
        _add_colored_run(
            ans_para, "Answer:  ", bold=True, color=COLOR_LABEL, size_pt=11
        )
        ans_para.add_run(str(q.get("answer", "")))

        # ── Explanation ────────────────────────────────────────────────────
        if q.get("explanation"):
            exp_para = doc.add_paragraph()
            _add_colored_run(
                exp_para, "Explanation:  ", bold=True, color=COLOR_LABEL, size_pt=11
            )
            exp_para.add_run(q["explanation"])

        # ── Concepts tested (scenario-based) ──────────────────────────────
        if q.get("concepts_tested"):
            cpt_para = doc.add_paragraph()
            _add_colored_run(
                cpt_para, "Concepts tested:  ", bold=True, color=COLOR_LABEL, size_pt=11
            )
            cpt_para.add_run(", ".join(q["concepts_tested"]))

        # ── Divider ────────────────────────────────────────────────────────
        doc.add_paragraph("─" * 60)
