import os
from pathlib import Path
import tempfile
from docx import Document

from src.export.docx import DOCXExporter


def test_docx_exporter_creates_file_and_content():
    exporter = DOCXExporter()

    questions = [
        {
            "question": "What is Python?",
            "options": ["A snake", "A programming language", "A car"],
            "answer": "A programming language",
            "explanation": "Python is a programming language.",
            "concepts_tested": ["Programming", "Basics"],
        }
    ]

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, "output.docx")

        result = exporter.export(questions, output_path)

        assert os.path.exists(result)

        doc = Document(result)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        assert "Exam Question Bank" in full_text
        assert "Question 1:" in full_text
        assert "What is Python?" in full_text
        assert "Options:" in full_text
        assert "Answer:" in full_text
        assert "Explanation:" in full_text
        assert "Concepts tested:" in full_text


def test_docx_exporter_multiple_questions():
    exporter = DOCXExporter()

    questions = [
        {"question": "Q1", "answer": "A"},
        {"question": "Q2", "answer": "B"},
    ]

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, "multi.docx")

        result = exporter.export(questions, output_path)

        assert os.path.exists(result)

        doc = Document(result)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        assert "Question 1:" in full_text
        assert "Question 2:" in full_text
        assert "Q1" in full_text
        assert "Q2" in full_text


def test_docx_exporter_question_without_options():
    exporter = DOCXExporter()

    questions = [{"question": "Is Python interpreted?", "answer": "True"}]

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, "no_options.docx")

        result = exporter.export(questions, output_path)

        assert os.path.exists(result)

        doc = Document(result)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        assert "Is Python interpreted?" in full_text
        assert "Answer:" in full_text


def test_docx_exporter_handles_empty_questions():
    exporter = DOCXExporter()

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, "empty.docx")

        result = exporter.export([], output_path)

        assert os.path.exists(result)

        doc = Document(result)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        assert "Total questions: 0" in full_text


def test_docx_file_is_not_empty(tmp_path):
    exporter = DOCXExporter()

    questions = [{"question": "Sample?", "answer": "Yes"}]

    output_file = tmp_path / "file.docx"

    result = exporter.export(questions, str(output_file))

    assert os.path.exists(result)
    assert Path(result).stat().st_size > 0
